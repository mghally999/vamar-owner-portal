"""
Builds vamar-import-template.xlsx: the sheet Vamar fills in to load its real data through the
dashboard's Import page.

  python tooling/golive/make_import_template.py

Writes apps/dashboard/public/vamar-import-template.xlsx (downloadable from the Import page).
Each data sheet's headers start with the field key, so the Import page maps every column by
itself. Lists, number ranges and lengths are enforced by Excel data validation; the dashboard
validates again before anything is saved.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "apps" / "dashboard" / "public" / "vamar-import-template.xlsx"

SLATE = "25323A"
MIST = "E8E8E8"
HEADER_FONT = Font(name="Almarai", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=SLATE)
NOTE_FONT = Font(name="Almarai", italic=True, color="4C575E", size=10)
BODY_FONT = Font(name="Almarai")
UNIT_TYPES = ["apartment", "villa", "shop", "office", "warehouse", "other"]
ROWS = 5000  # validation applies to this many rows

SHEETS = {
    "Owners": [
        ("nameAr", "الاسم بالعربي", "مطلوب", 28),
        ("nameEn", "Name in English", "مطلوب", 28),
        ("phoneE164", "الجوال، مثل 0501234567 أو +971501234567", "مطلوب، رقم الدخول للتطبيق", 26),
        ("email", "البريد الإلكتروني", "اختياري", 28),
        ("preferredLang", "لغة التطبيق: ar أو en", "اختياري، ar إذا ترك فارغاً", 18),
    ],
    "Buildings": [
        ("nameAr", "اسم المبنى بالعربي", "مطلوب، يكتب في ورقة الوحدات كما هو", 28),
        ("nameEn", "Building name in English", "مطلوب", 28),
        ("area", "المنطقة، مثل ديرة", "مطلوب", 20),
        ("address", "العنوان", "اختياري", 36),
    ],
    "Units": [
        ("building", "اسم المبنى كما في ورقة المباني (عربي أو إنجليزي)", "مطلوب", 30),
        ("unitNo", "رقم الوحدة، مثل 1204 أو G-02", "مطلوب، لا يتكرر في نفس المبنى", 18),
        ("type", "النوع: apartment / villa / shop / office / warehouse / other", "اختياري، apartment إذا ترك فارغاً", 22),
        ("bedrooms", "عدد غرف النوم (0 إلى 20)", "اختياري", 16),
        ("sizeSqft", "المساحة بالقدم المربع", "اختياري", 16),
    ],
    "Tenants": [
        ("name", "اسم المستأجر أو الشركة", "مطلوب", 30),
        ("phone", "الجوال", "اختياري", 20),
        ("email", "البريد الإلكتروني", "اختياري", 28),
        ("company", "اسم الشركة إن وجد", "اختياري", 28),
    ],
    "Shares": [
        ("building", "اسم المبنى كما في ورقة المباني", "مطلوب", 28),
        ("unitNo", "رقم الوحدة كما في ورقة الوحدات", "مطلوب", 16),
        ("ownerPhone", "جوال المالك كما في ورقة الملاك", "مطلوب", 22),
        ("sharePct", "حصة المالك بالنسبة المئوية، مثل 100 أو 60 أو 33.33", "مطلوب؛ حصص الوحدة الواحدة مجموعها 100", 18),
        ("fromDate", "تاريخ بداية الملكية YYYY-MM-DD", "اختياري، تاريخ اليوم إذا ترك فارغاً", 18),
    ],
    "Contracts": [
        ("building", "اسم المبنى كما في ورقة المباني", "مطلوب", 28),
        ("unitNo", "رقم الوحدة", "مطلوب", 16),
        ("tenantName", "اسم المستأجر كما في ورقة المستأجرين", "مطلوب", 28),
        ("tenantPhone", "جوال المستأجر", "اختياري، يلزم إذا تكرر اسم المستأجر", 20),
        ("startDate", "تاريخ بداية العقد YYYY-MM-DD", "مطلوب", 16),
        ("endDate", "تاريخ نهاية العقد YYYY-MM-DD", "مطلوب، بعد البداية", 16),
        ("annualRentAed", "الإيجار السنوي بالدرهم، مثل 63000", "مطلوب، منزلتان عشريتان كحد أقصى", 18),
        ("chequeCount", "عدد الشيكات (1 إلى 12)", "مطلوب", 14),
        ("ejariNo", "رقم إيجاري", "اختياري", 18),
    ],
    "Cheques": [
        ("building", "اسم المبنى كما في ورقة المباني", "مطلوب", 28),
        ("unitNo", "رقم الوحدة", "مطلوب", 16),
        ("contractStart", "تاريخ بداية العقد الذي يتبعه الشيك YYYY-MM-DD", "مطلوب، كما في ورقة العقود", 18),
        ("seq", "ترتيب الشيك في العقد (1، 2، 3 ...)", "مطلوب، من 1 إلى عدد شيكات العقد", 12),
        ("dueDate", "تاريخ استحقاق الشيك YYYY-MM-DD", "مطلوب، ضمن مدة العقد", 16),
        ("amountAed", "مبلغ الشيك بالدرهم، مثل 15750", "مطلوب، منزلتان عشريتان كحد أقصى", 16),
        ("chequeNo", "رقم الشيك", "مطلوب", 16),
        ("bankName", "اسم البنك", "مطلوب", 22),
    ],
}

EXAMPLES = {
    "Owners": ["خالد المنصوري", "Khalid Al Mansoori", "+971501234567", "khalid@example.com", "ar"],
    "Buildings": ["برج النور", "Al Noor Tower", "ديرة", "شارع المطار، دبي"],
    "Units": ["برج النور", "1204", "apartment", 2, 1150],
    "Tenants": ["سارة خليل", "+971559876543", "", ""],
    "Shares": ["برج النور", "1204", "+971501234567", 100, "2026-01-01"],
    "Contracts": ["برج النور", "1204", "سارة خليل", "+971559876543", "2026-01-01", "2026-12-31", 63000, 4, ""],
    "Cheques": ["برج النور", "1204", "2026-01-01", 1, "2026-01-01", 15750, "836466", "Emirates NBD"],
}

DATE_COLUMNS = {"fromDate", "startDate", "endDate", "contractStart", "dueDate"}


def instructions(ws):
    ws.sheet_view.rightToLeft = True
    ws.column_dimensions["A"].width = 110
    lines = [
        ("نموذج استيراد بيانات فامار", Font(name="Almarai", bold=True, size=16, color=SLATE)),
        ("", None),
        ("طريقة الاستخدام", Font(name="Almarai", bold=True, size=13, color=SLATE)),
        ("1. املأ الأوراق واستوردها بالترتيب: الملاك، المباني، الوحدات، المستأجرون، الحصص، العقود، ثم الشيكات. صف واحد لكل سجل، ابتداءً من الصف الثاني.", BODY_FONT),
        ("2. لا تغيّر الصف الأول (أسماء الأعمدة) ولا تحذفه. مرّر المؤشر على اسم العمود لرؤية شرحه، والشرح نفسه في الأسفل.", BODY_FONT),
        ("3. في لوحة التحكم افتح «استيراد»، اختر الملف، ثم اختر الورقة ونوع البيانات. تُطابق الأعمدة تلقائياً.", BODY_FONT),
        ("4. راجع نتيجة الفحص: الصفوف السليمة تُستورد، والصفوف التي فيها أخطاء تظهر مع السبب لتصحيحها ثم رفعها مرة أخرى.", BODY_FONT),
        ("5. رفع نفس الملف مرتين لا يكرر البيانات: يظهر لك نتيجة الاستيراد السابق.", BODY_FONT),
        ("", None),
        ("قواعد مهمة", Font(name="Almarai", bold=True, size=13, color=SLATE)),
        ("• رقم الجوال هو رقم دخول المالك إلى التطبيق. أرقام الإمارات تقبل بأي صيغة: 0501234567 أو 971501234567 أو +971501234567.", BODY_FONT),
        ("• اسم المبنى في ورقة الوحدات يجب أن يطابق اسمه في ورقة المباني (العربي أو الإنجليزي).", BODY_FONT),
        ("• الحصص: صف لكل مالك في الوحدة، ومجموع حصص الوحدة الواحدة 100 بالضبط. الوحدة التي لها ملاك من قبل لا تتغير من الاستيراد (تُعدّل من شاشة الوحدة).", BODY_FONT),
        ("• العقود: لا يتداخل عقدان على نفس الوحدة. المستأجر يجب أن يكون في ورقة المستأجرين. يُنشأ العقد مسودة إلى أن تصل شيكاته.", BODY_FONT),
        ("• الشيكات: كل شيك يتبع عقداً موجوداً (نفس المبنى والوحدة وتاريخ بداية العقد). عندما يكتمل آخر شيك ومجموع الشيكات يساوي الإيجار السنوي يصبح العقد سارياً تلقائياً؛ إن اختلف المجموع يبقى مسودة لتفعيله من لوحة التحكم مع ذكر السبب.", BODY_FONT),
        ("• المبالغ بالدرهم (مثل 63000 أو 15750.50) وتُحفظ بالفلس. التواريخ بصيغة YYYY-MM-DD أو تاريخ Excel.", BODY_FONT),
        ("• حالة الشيكات (مودَع، تم الصرف، مرتجع) تُحدّث من لوحة التحكم بعد الاستيراد (تحديث جماعي في شاشة الشيكات).", BODY_FONT),
        ("• حتى 5,000 صف في كل ورقة. الملف الأكبر يقسّم إلى ملفات.", BODY_FONT),
        ("", None),
        ("شرح الأعمدة", Font(name="Almarai", bold=True, size=13, color=SLATE)),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        if font:
            cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    row = len(lines) + 1
    for name, cols in SHEETS.items():
        ws.cell(row=row, column=1, value=f"{name}: " + " · ".join(
            f"{key}: {hint} ({rule})" for key, hint, rule, _w in cols
        )).alignment = Alignment(wrap_text=True)
        row += 1
    row += 1
    ws.cell(row=row, column=1, value="أمثلة (لا تنسخها إلى الأوراق)").font = Font(name="Almarai", bold=True, size=13, color=SLATE)
    row += 1
    for name, cols in SHEETS.items():
        ws.cell(row=row, column=1, value=f"{name}: " + " | ".join(
            f"{key}={value}" for (key, *_), value in zip(cols, EXAMPLES[name])
        )).font = NOTE_FONT
        row += 1


def data_sheet(wb, name, cols):
    ws = wb.create_sheet(name)
    ws.freeze_panes = "A2"
    for c, (key, hint, rule, width) in enumerate(cols, start=1):
        head = ws.cell(row=1, column=c, value=key)
        head.font = HEADER_FONT
        head.fill = HEADER_FILL
        head.alignment = Alignment(horizontal="center")
        head.comment = Comment(f"{hint}\n{rule}", "Vamar", width=260, height=90)
        ws.column_dimensions[head.column_letter].width = width
    last = ROWS + 1

    def col(key):
        for c, (k, *_rest) in enumerate(cols, start=1):
            if k == key:
                return ws.cell(row=1, column=c).column_letter
        raise KeyError(key)

    def add(dv, key):
        ws.add_data_validation(dv)
        dv.add(f"{col(key)}2:{col(key)}{last}")

    required_text = lambda title: DataValidation(  # noqa: E731
        type="textLength", operator="between", formula1="1", formula2="200",
        allow_blank=True, showErrorMessage=True, errorTitle=title,
        error="النص طويل جداً (200 حرف كحد أقصى).",
    )
    if name == "Owners":
        add(required_text("الاسم"), "nameAr")
        add(required_text("Name"), "nameEn")
        add(DataValidation(type="textLength", operator="between", formula1="9", formula2="16",
                           allow_blank=True, showErrorMessage=True, errorTitle="الجوال",
                           error="رقم الجوال بين 9 و16 خانة، مثل 0501234567."), "phoneE164")
        add(DataValidation(type="list", formula1='"ar,en"', allow_blank=True,
                           showErrorMessage=True, errorTitle="اللغة", error="اختر ar أو en."),
            "preferredLang")
    if name == "Buildings":
        add(required_text("الاسم"), "nameAr")
        add(required_text("Name"), "nameEn")
    if name == "Units":
        add(DataValidation(type="list", formula1='"' + ",".join(UNIT_TYPES) + '"', allow_blank=True,
                           showErrorMessage=True, errorTitle="النوع", error="اختر نوعاً من القائمة."),
            "type")
        add(DataValidation(type="whole", operator="between", formula1="0", formula2="20",
                           allow_blank=True, showErrorMessage=True, errorTitle="غرف النوم",
                           error="رقم صحيح من 0 إلى 20."), "bedrooms")
        add(DataValidation(type="whole", operator="between", formula1="1", formula2="1000000",
                           allow_blank=True, showErrorMessage=True, errorTitle="المساحة",
                           error="رقم صحيح أكبر من صفر."), "sizeSqft")
    if name == "Tenants":
        add(required_text("الاسم"), "name")
    if name in ("Shares", "Contracts", "Cheques"):
        add(required_text("المبنى"), "building")
    if name == "Shares":
        add(DataValidation(type="decimal", operator="between", formula1="0.01", formula2="100",
                           allow_blank=True, showErrorMessage=True, errorTitle="الحصة",
                           error="نسبة من 0.01 إلى 100."), "sharePct")
    if name == "Contracts":
        add(required_text("المستأجر"), "tenantName")
        add(DataValidation(type="decimal", operator="greaterThan", formula1="0",
                           allow_blank=True, showErrorMessage=True, errorTitle="الإيجار",
                           error="مبلغ بالدرهم أكبر من صفر."), "annualRentAed")
        add(DataValidation(type="whole", operator="between", formula1="1", formula2="12",
                           allow_blank=True, showErrorMessage=True, errorTitle="عدد الشيكات",
                           error="رقم صحيح من 1 إلى 12."), "chequeCount")
    if name == "Cheques":
        add(DataValidation(type="whole", operator="between", formula1="1", formula2="12",
                           allow_blank=True, showErrorMessage=True, errorTitle="ترتيب الشيك",
                           error="رقم صحيح من 1 إلى 12."), "seq")
        add(DataValidation(type="decimal", operator="greaterThan", formula1="0",
                           allow_blank=True, showErrorMessage=True, errorTitle="المبلغ",
                           error="مبلغ بالدرهم أكبر من صفر."), "amountAed")
    for key, *_rest in cols:
        if key in DATE_COLUMNS:
            add(DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)",
                               formula2="DATE(2100,12,31)", allow_blank=True, showErrorMessage=True,
                               errorTitle="التاريخ", error="تاريخ صحيح بصيغة YYYY-MM-DD."), key)
            for r in range(2, 202):
                ws[f"{col(key)}{r}"].number_format = "yyyy-mm-dd"


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "اقرأني"
    instructions(ws)
    for name, cols in SHEETS.items():
        data_sheet(wb, name, cols)
    wb.properties.creator = "Vamar Real Estate"
    wb.properties.title = "Vamar import template"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
